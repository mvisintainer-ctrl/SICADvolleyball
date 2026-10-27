import datetime
import html
import os
from pathlib import Path
import sys

from PySide2.QtCore import Qt, QStandardPaths, QByteArray, QUrl
from PySide2.QtGui import QColor, QFont, QPixmap, QTextDocument, QPainter, QImage
from PySide2.QtPrintSupport import QPrinter, QPrintDialog
from PySide2.QtWidgets import (QApplication, QMainWindow, QTableWidget, QTableWidgetItem,
    QTabWidget, QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QFileDialog,
    QMessageBox, QAction, QInputDialog, QDialog, QLineEdit, QComboBox, QLabel, QAbstractItemView)
from openpyxl.utils import get_column_letter, coordinate_to_tuple

try:
    from .model import Store, encode
    from .drawings import Drawings
except ImportError:
    from model import Store, encode
    from drawings import Drawings

ROOT = Path(getattr(sys, '_MEIPASS', Path(__file__).resolve().parents[1]))


def color(value, fallback):
    if value and value.type == 'rgb' and isinstance(value.rgb, str):
        return '#' + value.rgb[-6:]
    return fallback


class Sheet(QTableWidget):
    def __init__(self, sheet, store):
        super().__init__(sheet.max_row, sheet.max_column)
        self.sheet, self.store = sheet, store
        self.setHorizontalHeaderLabels([get_column_letter(c + 1) for c in range(sheet.max_column)])
        self.verticalHeader().setDefaultSectionSize(22)
        self.setAlternatingRowColors(False)
        self.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.setShowGrid(False)
        self.horizontalHeader().hide()
        self.verticalHeader().hide()
        self.setStyleSheet('QTableWidget {background:white; color:#203764; border:0;} QTableWidget::item:selected {background:#ccefff;color:#203764;}')
        self.loading = True
        overrides = store.overrides(sheet.title)
        for c in range(1, sheet.max_column + 1):
            dim = next((d for d in sheet.column_dimensions.values() if d.min <= c <= d.max), None)
            self.setColumnWidth(c - 1, max(12, int((dim.width if dim else 8.43) * 7)))
        for row, dim in sheet.row_dimensions.items():
            if dim.height:
                self.setRowHeight(row - 1, max(12, int(dim.height * 1.33)))
            self.setRowHidden(row - 1, bool(dim.hidden))
        for row in sheet:
            for cell in row:
                if cell.__class__.__name__ == 'MergedCell':
                    continue
                value = overrides.get(cell.coordinate, encode(cell.value))
                if cell.data_type == 'f' and cell.coordinate not in overrides:
                    value = datetime.date.today().strftime('%d/%m/%Y') if 'TODAY()' in str(cell.value) else str(len(store.records()))
                item = QTableWidgetItem('' if value is None else str(value))
                item.setFont(QFont(cell.font.name or 'Calibri', int(cell.font.sz or 11), QFont.Bold if cell.font.bold else QFont.Normal))
                item.setForeground(QColor(color(cell.font.color, '#203764')))
                if cell.fill.patternType == 'solid':
                    item.setBackground(QColor(color(cell.fill.fgColor, '#ffffff')))
                alignment = {'center': Qt.AlignHCenter, 'right': Qt.AlignRight}.get(cell.alignment.horizontal, Qt.AlignLeft)
                item.setTextAlignment(alignment | Qt.AlignVCenter)
                self.setItem(cell.row - 1, cell.column - 1, item)
        for merged in sheet.merged_cells.ranges:
            self.setSpan(merged.min_row - 1, merged.min_col - 1, merged.max_row - merged.min_row + 1, merged.max_col - merged.min_col + 1)
        self.loading = False
        self.cellChanged.connect(self.changed)

    def paintEvent(self, event):
        super().paintEvent(event)
        # Excel lets right-aligned labels overflow into empty cells on their left.
        painter = QPainter(self.viewport())
        first = max(0,self.rowAt(0))
        last = self.rowAt(self.viewport().height()-1)
        if last < 0: last = min(self.rowCount()-1,first+100)
        for row in range(first,last+1):
            for col in range(self.columnCount()):
                item = self.item(row,col)
                if not item or not item.text() or not item.textAlignment() & Qt.AlignRight:
                    continue
                rect = self.visualItemRect(item)
                painter.setFont(item.font())
                width = painter.fontMetrics().horizontalAdvance(item.text())+8
                if width <= rect.width(): continue
                rect.setLeft(rect.right()-width)
                painter.fillRect(rect,item.background() if item.background().style()!=Qt.NoBrush else QColor('white'))
                painter.setPen(item.foreground().color())
                painter.drawText(rect,Qt.AlignRight|Qt.AlignVCenter,item.text())
        painter.end()

    def text(self, address):
        row, col = coordinate_to_tuple(address)
        item = self.item(row - 1, col - 1)
        return item.text() if item else ''

    def put(self, address, value):
        row, col = coordinate_to_tuple(address)
        if self.item(row - 1, col - 1) is None:
            self.setItem(row - 1, col - 1, QTableWidgetItem())
        self.item(row - 1, col - 1).setText('' if value is None else str(value))

    def changed(self, row, col):
        if not self.loading:
            self.store.set_cell(self.sheet.title, f'{get_column_letter(col + 1)}{row + 1}', self.item(row, col).text())


class Window(QMainWindow):
    def __init__(self, store):
        super().__init__()
        self.store, self.photo = store, None
        self.current_code = None
        self.setWindowTitle('SICAD Volleyball — Versão de teste')
        self.resize(1450, 900)
        central = QWidget()
        self.setCentralWidget(central)
        layout = QVBoxLayout(central)
        self.tabs = QTabWidget()
        self.tabs.setStyleSheet('QTabBar::tab {background:#e81c06;color:white;padding:12px 24px;} QTabBar::tab:selected {background:#00b0f0;color:#203764;}')
        self.sheets = {}
        for sheet in store.workbook:
            widget = Sheet(sheet, store)
            self.sheets[sheet.title] = widget
            self.tabs.addTab(widget, 'Pagamentos' if sheet.title == 'Config' else sheet.title)
        actions = QHBoxLayout()
        for title, callback in [('Novo atleta', self.new), ('Cadastrar', lambda: self.save(True)),
                ('Alterar', lambda: self.save(False)), ('Excluir', self.delete), ('Limpar', self.clear),
                ('Anterior', lambda: self.move(-1)), ('Próximo', lambda: self.move(1)),
                ('Pesquisar', self.search), ('Foto', self.choose_photo), ('Remover foto', self.remove_photo),
                ('Salvar PDF', self.pdf), ('Imprimir ficha', self.print_sheet)]:
            button = QPushButton(title)
            button.clicked.connect(callback)
            actions.addWidget(button)
        layout.addLayout(actions)
        categories = QHBoxLayout()
        for title, start, end, sheet in [('Pessoais',11,27,'Cadastro'), ('Adicionais',28,43,'Cadastro'),
                ('Formação',44,58,'Cadastro'), ('Mirim',12,36,'Config'), ('Infantil',37,61,'Config'), ('Infanto-juvenil',62,86,'Config')]:
            button = QPushButton(title)
            button.clicked.connect(lambda checked=False, a=start, b=end, s=sheet: self.section(s,a,b))
            categories.addWidget(button)
        layout.addLayout(categories)
        layout.addWidget(self.tabs)
        self.photo_label = QLabel('Foto do atleta: nenhuma imagem selecionada')
        self.photo_label.setMaximumHeight(140)
        layout.addWidget(self.photo_label)
        self.statusBar().showMessage('Alterações nas células são salvas localmente. Use Cadastrar/Alterar para gravar a ficha no cadastro de atletas.')
        file_menu = self.menuBar().addMenu('Arquivo')
        for title, callback in [('Backup dos dados', self.backup), ('Mostrar todas as linhas', self.show_rows)]:
            action = QAction(title, self)
            action.triggered.connect(callback)
            file_menu.addAction(action)
        self.drawings = [Drawings(self.sheets[name], ROOT/'assets'/'template.xlsx', number, self) for number,name in enumerate(self.sheets,1)]
        for address in store.cad:
            row,col=coordinate_to_tuple(address)
            item=self.sheets['Cadastro'].item(row-1,col-1)
            if item:
                item.setBackground(QColor('#e9f6fc'))
                item.setToolTip('Clique duas vezes para editar. Grave com Cadastrar ou Alterar.')
        self.refresh_data()
        self.sheets['Dados'].cellChanged.connect(self.edit_data)
        self.clear()

    def section(self, name, start, end):
        self.tabs.setCurrentWidget(self.sheets[name])
        first, last = (11,58) if name == 'Cadastro' else (12,86)
        for row in range(first, last + 1):
            self.sheets[name].setRowHidden(row - 1, not start <= row <= end)
        for drawing in self.drawings:
            drawing.position()

    def show_rows(self):
        for sheet in self.sheets.values():
            for row in range(sheet.rowCount()):
                sheet.setRowHidden(row, False)
        for drawing in self.drawings:
            drawing.position()

    def run_macro(self, name):
        actions = {
            'NovoAtleta': self.new, 'Cadastrar': lambda: self.save(True),
            'Alterar': lambda: self.save(False), 'Excluir': self.delete,
            'Cancelar': self.clear, 'BuscarFoto': self.choose_photo,
            'CancelarFoto': self.remove_photo, 'AbrirPesquisa': self.search,
            'Anterior': lambda: self.move(-1), 'Proximo': lambda: self.move(1),
            'BuscarCodigo': lambda: self.load(self.sheets['Cadastro'].text(self.store.cad[0])),
            'GerarPDF': self.pdf, 'ImprimirFicha': self.print_sheet,
            'menuIP': lambda: self.section('Cadastro',11,27),
            'menuIA': lambda: self.section('Cadastro',28,43),
            'menuIF': lambda: self.section('Cadastro',44,58),
            'menuMirim': lambda: self.section('Config',12,36),
            'menuInfantil': lambda: self.section('Config',37,61),
            'menuInfanto': lambda: self.section('Config',62,86),
            'AtivarDesativarTela': lambda: self.showNormal() if self.isFullScreen() else self.showFullScreen(),
            'TelaMaxima': self.showFullScreen, 'TelaNormal': self.showNormal,
        }
        for macro, sheet in [('mCadastro','Cadastro'),('mFicha','Ficha'),('mDados','Dados'),('mConfig','Config')]:
            actions[macro] = lambda s=sheet: self.tabs.setCurrentWidget(self.sheets[s])
        if name in actions:
            actions[name]()
        else:
            QMessageBox.information(self,'Ação',f'Ação ainda não convertida: {name}')

    def values(self):
        return [self.sheets['Cadastro'].text(address) for address in self.store.cad]

    def display_photo(self):
        if self.photo:
            pixmap = QPixmap()
            if pixmap.loadFromData(QByteArray(self.photo)):
                self.photo_label.setPixmap(pixmap.scaled(140,140,Qt.KeepAspectRatio,Qt.SmoothTransformation))
                return
        self.photo_label.setText('Foto do atleta: nenhuma imagem selecionada')

    def clear(self):
        for address in self.store.cad:
            self.sheets['Cadastro'].put(address, '')
        self.photo, self.current_code = None, None
        self.display_photo()

    def new(self):
        self.clear()
        self.sheets['Cadastro'].put(self.store.cad[0], self.store.next_code())
        self.section('Cadastro',11,27)

    def save(self, creating):
        data = self.values()
        if not creating and self.current_code is not None and str(data[0]) != str(self.current_code):
            QMessageBox.warning(self, 'Código', 'O código identifica o atleta e não pode ser alterado. Crie um novo cadastro para outro código.')
            return
        try:
            self.store.save(data, self.photo, creating)
        except ValueError as error:
            QMessageBox.warning(self, 'Cadastro', str(error))
            return
        self.current_code = str(data[0])
        self.refresh_data()
        self.load(self.current_code)
        self.statusBar().showMessage('Cadastro salvo com sucesso.')

    def refresh_data(self):
        sheet = self.sheets['Dados']
        sheet.loading = True
        sheet.setRowCount(max(sheet.rowCount(), len(self.store.records()) + 11))
        for row in range(10, sheet.rowCount()):
            for col in range(38):
                item = sheet.item(row,col)
                if item:
                    item.setText('')
        self.data_codes = {}
        for row, data in enumerate(self.store.records(), 10):
            self.data_codes[row] = str(data[0])
            for col, value in enumerate(data):
                address = f'{get_column_letter(col+1)}{row+1}'
                sheet.put(address,value)
        sheet.loading = False
        self.sheets['Config'].put('A10',len(self.store.records()))

    def edit_data(self, row, col):
        sheet = self.sheets['Dados']
        if sheet.loading or row not in self.data_codes or col >= 37:
            return
        old_code = self.data_codes[row]
        try:
            new_code = self.store.update_field(old_code,col,sheet.item(row,col).text())
        except ValueError as error:
            QMessageBox.warning(self,'Editar cadastro',str(error))
            self.refresh_data()
            return
        self.data_codes[row] = new_code
        if self.current_code == old_code:
            self.load(new_code)
        self.statusBar().showMessage('Cadastro atualizado pelo banco de dados.')

    def load(self, code):
        result = self.store.get(code)
        if not result:
            return
        data, self.photo = result
        self.current_code = str(data[0])
        for i,value in enumerate(data):
            self.sheets['Cadastro'].put(self.store.cad[i],value)
            self.sheets['Ficha'].put(self.store.ficha[i],value)
        self.display_photo()

    def delete(self):
        code = self.sheets['Cadastro'].text(self.store.cad[0])
        if not self.store.get(code):
            QMessageBox.warning(self,'Excluir','Selecione um cadastro existente.')
            return
        if QMessageBox.question(self,'Excluir','Excluir este atleta? Faça um backup antes se precisar recuperar o cadastro.') == QMessageBox.Yes:
            self.store.delete(code)
            self.clear()
            for address in self.store.ficha:
                self.sheets['Ficha'].put(address,'')
            self.refresh_data()

    def move(self, delta):
        codes = [str(row[0]) for row in self.store.records()]
        if not codes:
            return
        index = codes.index(self.current_code) if self.current_code in codes else (-1 if delta > 0 else len(codes))
        self.load(codes[max(0,min(len(codes)-1,index+delta))])

    def search(self):
        dialog = QDialog(self)
        dialog.setWindowTitle('Pesquisar atleta')
        dialog.resize(800,500)
        layout = QVBoxLayout(dialog)
        query, field = QLineEdit(), QComboBox()
        query.setPlaceholderText('Digite para pesquisar…')
        field.addItems(['Nome','Categoria','Posição','Código'])
        layout.addWidget(field)
        layout.addWidget(query)
        results = QTableWidget(0,4)
        results.setHorizontalHeaderLabels(['Código','Nome','Categoria','Posição'])
        results.setSelectionBehavior(QAbstractItemView.SelectRows)
        results.setEditTriggers(QAbstractItemView.NoEditTriggers)
        results.setColumnWidth(1,350)
        layout.addWidget(results)
        count = QLabel()
        layout.addWidget(count)
        def filter_rows():
            column = [1,19,20,0][field.currentIndex()]
            records = [r for r in self.store.records() if query.text().casefold() in str(r[column] or '').casefold()]
            results.setRowCount(len(records))
            for i,row in enumerate(records):
                for j,col in enumerate([0,1,19,20]):
                    results.setItem(i,j,QTableWidgetItem(str(row[col] or '')))
            count.setText(f'{len(records)} atletas encontrados. Clique duas vezes para abrir.')
        def select(row,col):
            self.load(results.item(row,0).text())
            dialog.accept()
            if self.tabs.currentWidget() != self.sheets['Ficha']:
                self.section('Cadastro',11,27)
        query.textChanged.connect(filter_rows)
        field.currentIndexChanged.connect(filter_rows)
        results.cellDoubleClicked.connect(select)
        filter_rows()
        dialog.exec_()

    def choose_photo(self):
        path,_ = QFileDialog.getOpenFileName(self,'Selecionar foto','','Imagens (*.jpg *.jpeg *.png *.bmp)')
        if path:
            content = Path(path).read_bytes()
            pixmap = QPixmap()
            if not pixmap.loadFromData(content):
                QMessageBox.warning(self,'Foto','Não foi possível abrir a imagem.')
                return
            self.photo = content
            self.display_photo()

    def remove_photo(self):
        self.photo = None
        self.display_photo()

    def document(self):
        # Respect the original Ficha print area, merged cells, values and colours.
        sheet = self.sheets['Ficha']
        covered = set()
        spans = {}
        for merged in self.store.workbook['Ficha'].merged_cells.ranges:
            spans[(merged.min_row,merged.min_col)] = (merged.max_row-merged.min_row+1,merged.max_col-merged.min_col+1)
            for r in range(merged.min_row,merged.max_row+1):
                for c in range(merged.min_col,merged.max_col+1):
                    if (r,c)!=(merged.min_row,merged.min_col): covered.add((r,c))
        parts = ['<html><body><table width="100%" cellspacing="0" cellpadding="2">']
        for row in range(10,73):
            parts.append('<tr>')
            for col in range(10,43):
                if (row,col) in covered:
                    continue
                item = sheet.item(row-1,col-1)
                value = html.escape(item.text()) if item else ''
                rs,cs = spans.get((row,col),(1,1))
                bg = item.background().color().name() if item and item.background().style()!=Qt.NoBrush else '#ffffff'
                parts.append(f'<td rowspan="{rs}" colspan="{cs}" bgcolor="{bg}"><font color="#203764" size="2">{value}</font></td>')
            parts.append('</tr>')
        parts.append('</table>')
        if self.photo:
            parts.append('<p><img width="80" src="sicad-photo"></p>')
        parts.append('</body></html>')
        document = QTextDocument()
        document.setHtml(''.join(parts))
        if self.photo:
            document.addResource(QTextDocument.ImageResource,QUrl('sicad-photo'),QImage.fromData(self.photo))
        return document

    def pdf(self):
        if self.current_code is None:
            QMessageBox.warning(self,'Ficha','Pesquise e selecione um atleta antes de gerar a ficha.')
            return
        path,_ = QFileDialog.getSaveFileName(self,'Salvar ficha','Ficha.pdf','PDF (*.pdf)')
        if path:
            if not path.lower().endswith('.pdf'): path += '.pdf'
            printer = QPrinter(QPrinter.HighResolution)
            printer.setOutputFormat(QPrinter.PdfFormat)
            printer.setOutputFileName(path)
            self.document().print_(printer)
            self.statusBar().showMessage('PDF salvo em '+path)

    def print_sheet(self):
        if self.current_code is None:
            QMessageBox.warning(self,'Ficha','Selecione um atleta antes de imprimir.')
            return
        printer = QPrinter()
        if QPrintDialog(printer,self).exec_() == QDialog.Accepted:
            self.document().print_(printer)

    def backup(self):
        path,_ = QFileDialog.getSaveFileName(self,'Backup dos dados','SICAD-backup.sqlite3','SQLite (*.sqlite3)')
        if path:
            self.store.backup(path)
            self.statusBar().showMessage('Backup salvo em '+path)


def main():
    application = QApplication(sys.argv)
    application.setOrganizationName('SICAD')
    application.setApplicationName('SICADvolleyballTeste')
    folder = Path(QStandardPaths.writableLocation(QStandardPaths.AppLocalDataLocation))
    folder.mkdir(parents=True,exist_ok=True)
    store = Store(folder/'sicad.sqlite3', ROOT/'assets'/'template.xlsx')
    window = Window(store)
    window.show()
    sys.exit(application.exec_())


if __name__ == '__main__':
    main()
