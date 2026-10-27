"""Workbook import and local persistence. Original Excel file is never rewritten."""
import datetime
import json
import sqlite3
from pathlib import Path

import openpyxl


def encode(value):
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime('%d/%m/%Y')
    return value


class Store:
    def __init__(self, database, template):
        self.connection = sqlite3.connect(database)
        self.connection.executescript('''
            CREATE TABLE IF NOT EXISTS athlete (code TEXT PRIMARY KEY, data TEXT NOT NULL, photo BLOB);
            CREATE TABLE IF NOT EXISTS cell (sheet TEXT, address TEXT, value TEXT, PRIMARY KEY(sheet,address));
            CREATE TABLE IF NOT EXISTS metadata (key TEXT PRIMARY KEY, value TEXT);
        ''')
        self.workbook = openpyxl.load_workbook(template, data_only=False)
        self.fields = [self.workbook['Dados'].cell(10, i + 1).value or str(i) for i in range(37)]
        self.cad = [self._destination('Cad_' + str(i)) for i in range(37)]
        self.ficha = [self._destination('Ficha' + str(i)) for i in range(37)]
        if not self.connection.execute("SELECT 1 FROM metadata WHERE key='imported'").fetchone():
            # Import is one transaction, including the marker, so interrupted imports can retry.
            with self.connection:
                for row in self.workbook['Dados'].iter_rows(min_row=11, max_col=38):
                    if row[0].value is None:
                        continue
                    data = [encode(c.value) for c in row[:37]]
                    self.connection.execute('INSERT INTO athlete VALUES (?,?,NULL)', (str(data[0]), json.dumps(data, ensure_ascii=False)))
                self.connection.execute("INSERT INTO metadata VALUES ('imported','1')")

    def _destination(self, name):
        return next(self.workbook.defined_names[name].destinations)[1].replace('$', '')

    def records(self):
        rows = self.connection.execute('SELECT data FROM athlete ORDER BY rowid').fetchall()
        return [json.loads(row[0]) for row in rows]

    def get(self, code):
        row = self.connection.execute('SELECT data,photo FROM athlete WHERE code=?', (str(code),)).fetchone()
        return (json.loads(row[0]), row[1]) if row else None

    def save(self, data, photo=None, creating=False):
        if not str(data[0] or '').strip() or not str(data[1] or '').strip():
            raise ValueError('Informe o código e o nome do atleta.')
        code = str(data[0])
        exists = self.get(code)
        if creating and exists:
            raise ValueError('Este código já existe. Use Alterar para editar o cadastro.')
        if not creating and not exists:
            raise ValueError('Cadastro não encontrado. Use Cadastrar para um novo atleta.')
        with self.connection:
            self.connection.execute('INSERT INTO athlete VALUES (?,?,?) ON CONFLICT(code) DO UPDATE SET data=excluded.data, photo=excluded.photo', (code, json.dumps(data, ensure_ascii=False), photo))

    def delete(self, code):
        with self.connection:
            self.connection.execute('DELETE FROM athlete WHERE code=?', (str(code),))
        # Codes remain stable; deleting a record must not change other athletes' identities.

    def update_field(self, code, column, value):
        existing = self.get(code)
        if not existing:
            raise ValueError('Cadastro não encontrado.')
        data, photo = existing
        data[column] = value
        if column in (0, 1) and not str(value).strip():
            raise ValueError('Código e nome são obrigatórios.')
        new_code = str(data[0])
        if new_code != str(code) and self.get(new_code):
            raise ValueError('Este código já existe.')
        with self.connection:
            self.connection.execute('UPDATE athlete SET code=?,data=? WHERE code=?',
                (new_code,json.dumps(data,ensure_ascii=False),str(code)))
        return new_code

    def next_code(self):
        numbers = [int(str(r[0])) for r in self.records() if str(r[0]).isdigit()]
        return str(max(numbers, default=0) + 1)

    def set_cell(self, sheet, address, value):
        with self.connection:
            self.connection.execute('INSERT INTO cell VALUES (?,?,?) ON CONFLICT(sheet,address) DO UPDATE SET value=excluded.value', (sheet, address, value))

    def overrides(self, sheet):
        return dict(self.connection.execute('SELECT address,value FROM cell WHERE sheet=?', (sheet,)))

    def backup(self, path):
        with sqlite3.connect(path) as target:
            self.connection.backup(target)
